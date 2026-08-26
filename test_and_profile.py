import time
import os
import sys
import shutil
from pathlib import Path
import docx

# Ensure stdout handles UTF-8 on Windows terminal
sys.stdout.reconfigure(encoding='utf-8')

import rag_service

def run_tests_and_profile():
    print("=" * 60)
    print("STARTING PROFILE & TEST SUITE FOR DOCUMIND_RAG")
    print("=" * 60)
    
    pdf_path = Path("bcf34d2d-2a9e-499b-b678-612f7b6bd3f2.pdf")
    docx_path = Path("sample_test.docx")
    txt_path = Path("sample_test.txt")
    
    # Create sample docx and txt files if not exist
    if not docx_path.exists():
        doc = docx.Document()
        doc.add_heading('DocMind RAG Architecture', 0)
        doc.add_paragraph('This document describes the RAG pipeline built from scratch without orchestration frameworks.')
        doc.save(str(docx_path))
        
    if not txt_path.exists():
        txt_path.write_text("DocMind RAG provides document chunking, sentence transformer embeddings, and vector storage using ChromaDB.", encoding="utf-8")
        
    # --- 1. PROFILE INGESTION (PDF First Run vs Second Run Cache Hit) ---
    print("\n--- [1] Profiling Ingestion ---")
    
    # Clear cache dir to test first run
    if rag_service.CACHE_DIR.exists():
        shutil.rmtree(rag_service.CACHE_DIR)
    rag_service.CACHE_DIR.mkdir(exist_ok=True)
    
    t0 = time.perf_counter()
    ingested_pdf_1 = rag_service.ingest_document(pdf_path)
    t1 = time.perf_counter()
    first_ingest_time = t1 - t0
    print(f"First PDF Ingestion Time (Cold): {first_ingest_time:.4f}s | Chunks: {len(ingested_pdf_1['chunks'])}")
    
    t0 = time.perf_counter()
    ingested_pdf_2 = rag_service.ingest_document(pdf_path)
    t1 = time.perf_counter()
    second_ingest_time = t1 - t0
    print(f"Second PDF Ingestion Time (On-Disk Cache Hit): {second_ingest_time:.4f}s | Speedup: {first_ingest_time / max(second_ingest_time, 0.0001):.1f}x")
    
    # Test DOCX and TXT ingestion
    t0 = time.perf_counter()
    ingested_docx = rag_service.ingest_document(docx_path)
    t1 = time.perf_counter()
    print(f"DOCX Ingestion Time: {t1 - t0:.4f}s | Chunks: {len(ingested_docx['chunks'])}")
    
    t0 = time.perf_counter()
    ingested_txt = rag_service.ingest_document(txt_path)
    t1 = time.perf_counter()
    print(f"TXT Ingestion Time: {t1 - t0:.4f}s | Chunks: {len(ingested_txt['chunks'])}")
    
    # --- 2. PROFILING QUERY LATENCY & CACHING ---
    print("\n--- [2] Profiling Query Path ---")
    query = "What is the primary subject of this document?"
    
    t0 = time.perf_counter()
    ans, source = rag_service.answer_query(ingested_pdf_1, query, target_lang="English", stream=False)
    t1 = time.perf_counter()
    first_query_time = t1 - t0
    print(f"First Query Latency (Cold): {first_query_time:.4f}s | Source: {source}")
    print(f"Answer snippet: {ans[:150]}...")
    
    t0 = time.perf_counter()
    ans_cached, source_cached = rag_service.answer_query(ingested_pdf_1, query, target_lang="English", stream=False)
    t1 = time.perf_counter()
    cached_query_time = t1 - t0
    print(f"Second Query Latency (In-Memory LRU Cache Hit): {cached_query_time:.4f}s | Source: {source_cached}")
    print(f"Cache speedup: {first_query_time / max(cached_query_time, 0.00001):.1f}x")

    # --- 3. PROFILING RESPONSE STREAMING ---
    print("\n--- [3] Testing Response Streaming ---")
    t0 = time.perf_counter()
    stream_gen, stream_source = rag_service.answer_query(ingested_pdf_1, "Summarize the key points briefly.", target_lang="English", stream=True)
    ttft = None
    collected_chunks = []
    
    if hasattr(stream_gen, "__iter__") and not isinstance(stream_gen, str):
        for chunk in stream_gen:
            if ttft is None:
                ttft = time.perf_counter() - t0
            collected_chunks.append(chunk)
        total_stream_time = time.perf_counter() - t0
        print(f"Streaming Source: {stream_source}")
        print(f"Time-To-First-Token (TTFT): {ttft:.4f}s" if ttft else "TTFT: N/A")
        print(f"Total Streaming Generation Time: {total_stream_time:.4f}s | Output length: {len(''.join(collected_chunks))} chars")
    else:
        print(f"Non-streaming result returned: {stream_gen[:100]}...")

    # --- 4. TESTING MULTILINGUAL & TRANSLATION FALLBACK TILE ---
    print("\n--- [4] Testing Multilingual & Translation Fallback Chain ---")
    
    hindi_ans, hindi_source = rag_service.answer_query(ingested_pdf_1, "What is the document about?", target_lang="हिंदी (Hindi)")
    print(f"Hindi Answer Source: {hindi_source}")
    print(f"Hindi Answer: {hindi_ans[:150]}...")
    
    marathi_ans, marathi_source = rag_service.answer_query(ingested_pdf_1, "What is the document about?", target_lang="मराठी (Marathi)")
    print(f"Marathi Answer Source: {marathi_source}")
    print(f"Marathi Answer: {marathi_ans[:150]}...")

    # Test Translation Fallback when SARVAM_API_KEY is missing/unset
    old_sarvam_key = os.environ.get("SARVAM_API_KEY")
    if "SARVAM_API_KEY" in os.environ:
        del os.environ["SARVAM_API_KEY"]
        
    print("\nTesting Translation Tier 2/3 (Fallback when SARVAM_API_KEY is unset):")
    hi_fallback_ans, hi_fallback_source = rag_service.answer_query(ingested_docx, "What is described in this document?", target_lang="हिंदी (Hindi)")
    print(f"Fallback Translation Source: {hi_fallback_source}")
    print(f"Result: {hi_fallback_ans[:150]}...")

    # Restore env key
    if old_sarvam_key:
        os.environ["SARVAM_API_KEY"] = old_sarvam_key

    # --- 5. DOCUMENT SUMMARIZATION TESTS ---
    print("\n--- [5] Testing Document Summarization ---")
    summary_en, source_sum_en = rag_service.summarize_document(ingested_pdf_1, target_lang="English")
    print(f"English Summary Source: {source_sum_en} | Length: {len(summary_en)} chars")
    
    summary_hi, source_sum_hi = rag_service.summarize_document(ingested_pdf_1, target_lang="हिंदी (Hindi)")
    print(f"Hindi Summary Source: {source_sum_hi} | Length: {len(summary_hi)} chars")

    print("\n" + "=" * 60)
    print("TEST & PROFILE SUITE COMPLETED SUCCESSFULLY!")
    print("=" * 60)

if __name__ == "__main__":
    run_tests_and_profile()
